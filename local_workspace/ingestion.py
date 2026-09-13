import csv
import hashlib
import io
import json
import math
import re
from pathlib import Path

try:
    from . import ollama_client
except ImportError:
    import ollama_client

ALLOWED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx", ".csv", ".xlsx"}


def safe_name(name):
    base = Path(name or "document").name
    return re.sub(r"[^A-Za-z0-9._ -]+", "_", base)[:180]


def digest(data):
    return hashlib.sha256(data).hexdigest()


def extract(path):
    suffix = path.suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("Supported files: PDF, TXT, MD, DOCX, CSV and XLSX.")
    if suffix in {".txt", ".md"}:
        return [{"text": path.read_text(encoding="utf-8", errors="replace"), "location": "document"}]
    if suffix == ".pdf":
        from pypdf import PdfReader
        return [
            {"text": page.extract_text() or "", "location": f"page {number}"}
            for number, page in enumerate(PdfReader(str(path)).pages, 1)
        ]
    if suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for table_number, table in enumerate(doc.tables, 1):
            lines.append(f"Table {table_number}")
            lines.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return [{"text": "\n".join(lines), "location": "document"}]
    if suffix == ".csv":
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
        return [{"text": "\n".join(" | ".join(row) for row in rows), "location": "sheet CSV"}]
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    sections = []
    try:
        for sheet in workbook.worksheets:
            lines = []
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    lines.append(" | ".join(values))
            sections.append({"text": "\n".join(lines), "location": f"sheet {sheet.title}"})
    finally:
        workbook.close()
    return sections


def chunks(sections, filename, size=1100, overlap=180):
    output = []
    for section in sections:
        clean = re.sub(r"[ \t]+", " ", section["text"]).strip()
        start = 0
        index = 1
        while start < len(clean):
            end = min(len(clean), start + size)
            if end < len(clean):
                boundary = max(clean.rfind("\n", start, end), clean.rfind(". ", start, end))
                if boundary > start + size // 2:
                    end = boundary + 1
            content = clean[start:end].strip()
            if content:
                output.append({
                    "content": content,
                    "metadata": {"filename": filename, "location": section["location"], "chunk": index},
                })
                index += 1
            if end >= len(clean):
                break
            start = max(start + 1, end - overlap)
    return output


def cosine(left, right):
    if not left or not right or len(left) != len(right):
        return -1.0
    dot = sum(a * b for a, b in zip(left, right))
    denom = math.sqrt(sum(a * a for a in left)) * math.sqrt(sum(b * b for b in right))
    return dot / denom if denom else -1.0


def index_document(store, scope, filename, data, embedding_model="nomic-embed-text"):
    filename = safe_name(filename)
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("Supported files: PDF, TXT, MD, DOCX, CSV and XLSX.")
    sha = digest(data)
    existing = store.rows("SELECT * FROM documents WHERE scope=? AND sha256=?", (scope, sha))
    if existing:
        if existing[0]["status"] != "failed":
            return existing[0], True
        store.execute("DELETE FROM documents WHERE id=?", (existing[0]["id"],))
        (store.documents_dir / existing[0]["stored_name"]).unlink(missing_ok=True)
    stored_name = f"{sha[:16]}{suffix}"
    path = store.documents_dir / stored_name
    path.write_bytes(data)
    document_id, _ = store.execute(
        "INSERT INTO documents(scope,filename,stored_name,sha256,kind,status,created) VALUES(?,?,?,?,?,'indexing',?)",
        (scope, filename, stored_name, sha, suffix.lstrip("."), _utcnow()),
    )
    try:
        pieces = chunks(extract(path), filename)
        if not pieces:
            raise ValueError("No readable text was found in this document.")
        if len(pieces) > 5000:
            raise ValueError("This document is too large after extraction. Split it into smaller files.")
        vectors = []
        for start in range(0, len(pieces), 16):
            vectors.extend(ollama_client.embed([piece["content"] for piece in pieces[start:start + 16]], embedding_model))
        with store.connect() as con:
            con.executemany(
                "INSERT INTO chunks(document_id,scope,content,metadata,embedding) VALUES(?,?,?,?,?)",
                [(document_id, scope, piece["content"], json.dumps(piece["metadata"]), json.dumps(vector))
                 for piece, vector in zip(pieces, vectors)],
            )
            con.execute("UPDATE documents SET status='ready',chunk_count=? WHERE id=?", (len(pieces), document_id))
        store.audit(scope, "document_indexed", f"{filename}; {len(pieces)} chunks")
    except Exception:
        store.execute("UPDATE documents SET status='failed' WHERE id=?", (document_id,))
        raise
    return store.rows("SELECT * FROM documents WHERE id=?", (document_id,))[0], False


def _utcnow():
    from datetime import datetime, timezone
    return datetime.now(timezone.utc).isoformat()


def retrieve(store, scope, query, embedding_model="nomic-embed-text", limit=5):
    rows = store.rows("SELECT id,document_id,content,metadata,embedding FROM chunks WHERE scope=?", (scope,))
    if not rows:
        return []
    query_vector = ollama_client.embed(query, embedding_model)[0]
    ranked = []
    for row in rows:
        vector = store.decode_json(row["embedding"], [])
        ranked.append((cosine(query_vector, vector), row))
    output = []
    for score, row in sorted(ranked, key=lambda item: item[0], reverse=True)[:limit]:
        metadata = store.decode_json(row["metadata"], {})
        output.append({
            "id": row["id"], "document_id": row["document_id"], "content": row["content"],
            "filename": metadata.get("filename", "Document"),
            "location": metadata.get("location", "document"),
            "chunk": metadata.get("chunk"), "score": round(score, 4),
        })
    return output
