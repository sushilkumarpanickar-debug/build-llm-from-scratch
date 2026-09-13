"""Deterministic, local-only profiling for finance CSV and XLSX files."""
from __future__ import annotations

import csv
import io
import math
import re
from pathlib import Path

from openpyxl import load_workbook

SUPPORTED = {".csv", ".xlsx"}
MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_ROWS = 100_000
MONEY_HINTS = {
    "inflow": ("revenue", "sales", "receipt", "receipts", "collection", "collections", "credit", "inflow", "income"),
    "outflow": ("expense", "expenses", "cost", "payment", "payments", "debit", "outflow", "purchase"),
    "balance": ("balance", "closing", "net", "outstanding"),
}


def _number(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(float(value)) else None
    text = str(value).strip()
    if not text or not re.fullmatch(r"\(?[-+₹$€£]?\s*[\d,]+(?:\.\d+)?\)?", text):
        return None
    negative = text.startswith("(") and text.endswith(")")
    cleaned = re.sub(r"[^0-9.+-]", "", text)
    try:
        result = float(cleaned)
        return -result if negative else result
    except ValueError:
        return None


def _role(header):
    value = re.sub(r"[^a-z]+", " ", str(header).lower()).strip()
    for role, hints in MONEY_HINTS.items():
        if any(re.search(rf"\b{re.escape(hint)}\b", value) for hint in hints):
            return role
    return "numeric"


def _profile_sheet(name, rows):
    rows = list(rows)
    if not rows:
        return {"name": name, "rows": 0, "columns": 0, "numeric_columns": [], "warnings": ["Sheet is empty."]}
    width = max(len(row) for row in rows)
    header_row = list(rows[0]) + [None] * (width - len(rows[0]))
    headers = [str(value).strip() if value not in (None, "") else f"Column {index + 1}" for index, value in enumerate(header_row)]
    values = [[] for _ in range(width)]
    for row in rows[1:MAX_ROWS + 1]:
        padded = list(row) + [None] * (width - len(row))
        for index, value in enumerate(padded):
            number = _number(value)
            if number is not None:
                values[index].append(number)
    numeric = []
    for header, column in zip(headers, values):
        if not column:
            continue
        numeric.append({
            "name": header,
            "role": _role(header),
            "count": len(column),
            "sum": round(sum(column), 2),
            "average": round(sum(column) / len(column), 2),
            "minimum": round(min(column), 2),
            "maximum": round(max(column), 2),
        })
    warnings = []
    if len(rows) - 1 > MAX_ROWS:
        warnings.append(f"Only the first {MAX_ROWS:,} data rows were analysed.")
    if not numeric:
        warnings.append("No numeric columns were detected.")
    return {"name": name, "rows": max(0, min(len(rows) - 1, MAX_ROWS)), "columns": width, "headers": headers, "numeric_columns": numeric, "warnings": warnings}


def analyze_bytes(filename, data):
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError("Finance analysis supports CSV and XLSX files.")
    if not data or len(data) > MAX_FILE_BYTES:
        raise ValueError("Finance file must be between 1 byte and 25 MB.")
    if suffix == ".csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("CSV must use UTF-8 encoding.") from exc
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        sheets = [_profile_sheet("CSV", csv.reader(io.StringIO(text), dialect))]
    else:
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError(f"XLSX could not be read: {exc}") from exc
        try:
            sheets = [_profile_sheet(sheet.title, sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets]
        finally:
            workbook.close()
    roles = {"inflow": [], "outflow": [], "balance": []}
    for sheet in sheets:
        for column in sheet["numeric_columns"]:
            if column["role"] in roles:
                roles[column["role"]].append({"sheet": sheet["name"], "column": column["name"], "sum": column["sum"]})
    return {
        "filename": Path(filename).name,
        "size_bytes": len(data),
        "sheets": sheets,
        "recognized_finance_columns": roles,
        "method": "Deterministic local profiling; no AI-generated figures.",
    }


def analyze_path(path, allowed_root):
    root = Path(allowed_root).expanduser().resolve()
    candidate = Path(path).expanduser().resolve()
    try:
        candidate.relative_to(root)
    except ValueError as exc:
        raise ValueError("File is outside DAKSH_FINANCE_ROOT.") from exc
    if not candidate.is_file():
        raise ValueError("Finance file does not exist.")
    return analyze_bytes(candidate.name, candidate.read_bytes())
